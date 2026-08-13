"""Generates a real Bill of Materials file (.xlsx, via openpyxl) from a
CadAdapter.get_assembly_bom() result -- either flat (parts-only) or
indented (sub-assembly-grouped).

**Flat vs. indented is a real, deliberate distinction in manufacturing/
BOM practice, not just a formatting choice:**
- A **flat BOM** lists one row per unique (part, configuration) with its
  TOTAL quantity needed across the whole assembly, ignoring where in the
  tree it lives. This is what purchasing/stores actually needs: if a bolt
  is used both directly in the top assembly and inside a sub-assembly, the
  warehouse just needs to know "40 of these total", not which sub-assembly
  each one eventually ends up in.
- An **indented BOM** groups rows under the sub-assembly they belong to,
  preserving the assembly's build structure (SOLIDWORKS' own BOM table
  feature does the same thing). This is what an assembler/planner needs:
  it shows that a pin and bushing are built/kitted together as part of a
  Hinge Assembly sub-assembly, not loose parts that happen to share a
  final assembly -- which matters for staging work, sub-assembly-level
  make/buy decisions, and outsourcing a sub-assembly as its own unit.

Both are legitimate, standard BOM formats used for different purposes in
industry -- this module produces either, from the same rolled-up BOM data,
rather than picking one as "the" BOM format.
"""

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

_HEADER_FILL = PatternFill("solid", fgColor="1B2A41")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_SUBASSEMBLY_FILL = PatternFill("solid", fgColor="DDE3EC")
_SUBASSEMBLY_FONT = Font(bold=True, italic=True)

_COLUMNS = [
    "Item No",
    "Part Name",
    "Configuration",
    "Make/Buy",
    "Material",
    "Qty",
    "Unit Mass (kg)",
    "Total Mass (kg)",
    "Unit Cost (INR)",
    "Total Cost (INR)",
]


def _write_header(ws) -> None:
    ws.append(_COLUMNS)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _row_values(item_no: int, row: dict) -> list:
    return [
        item_no,
        row["part_name"],
        row["configuration"],
        row["classification"],
        row["material"] or "-",
        row["quantity_per_assembly"],
        row["unit_mass_kg"],
        row["total_mass_kg"],
        row["unit_cost_inr"],
        row["total_cost_inr"],
    ]


def _append_totals(ws, bom_result: dict) -> None:
    totals = bom_result["totals"]
    ws.append([])
    ws.append(["", "Unique parts", totals["unique_part_count"]])
    ws.append(["", "Total part instances", totals["total_instance_count"]])
    ws.append(["", "Total assembly mass (kg)", totals["total_assembly_mass_kg"]])
    ws.append(["", "Total cost -- 1 assembly (INR)", totals["total_assembly_cost_one_unit_inr"]])
    ws.append(
        [
            "",
            f"Total cost -- {bom_result['quantity']} assemblies (INR)",
            totals["total_assembly_cost_for_quantity_inr"],
        ]
    )
    if bom_result["missing_data"]:
        ws.append([])
        ws.append(["", "Missing data (excluded from totals above):"])
        for entry in bom_result["missing_data"]:
            ws.append(["", f"{entry['part_name']} ({entry['configuration']})", entry["reason"]])


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        length = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=10,
        )
        ws.column_dimensions[column_cells[0].column_letter].width = min(length + 2, 45)


def export_bom_flat(bom_result: dict, output_path: str) -> str:
    """Flat parts-only BOM: one row per unique (part, configuration), each
    with its total quantity across the whole assembly -- see module
    docstring for why this differs from the indented format.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM (Flat)"
    _write_header(ws)

    for item_no, row in enumerate(bom_result["bom"], start=1):
        ws.append(_row_values(item_no, row))

    _autosize(ws)
    _append_totals(ws, bom_result)
    wb.save(output_path)
    return output_path


def export_bom_indented(bom_result: dict, output_path: str) -> str:
    """Indented BOM: top-level parts first, then each sub-assembly's
    children grouped beneath a bold sub-assembly header row -- see module
    docstring for why this differs from the flat format.

    Groups by `parent_assembly` (AssemblyComponent.parent_assembly, first
    occurrence's parent if a part is shared across more than one
    sub-assembly -- see base_adapter.AssemblyComponent's docstring for
    that documented simplification). Only handles one level of indentation
    below the top assembly, matching this project's current traversal
    depth; a deeper multi-level assembly would need each nested
    sub-assembly's own `level` value walked recursively rather than a
    single parent-name grouping.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM (Indented)"
    _write_header(ws)

    rows = bom_result["bom"]
    top_level = [r for r in rows if r["parent_assembly"] is None]
    by_parent: dict[str, list] = {}
    for r in rows:
        if r["parent_assembly"] is not None:
            by_parent.setdefault(r["parent_assembly"], []).append(r)

    item_no = 1
    for row in top_level:
        ws.append(_row_values(item_no, row))
        item_no += 1

    for parent_name, children in by_parent.items():
        ws.append(["", parent_name, "", "Sub-Assembly", "", "", "", "", "", ""])
        for cell in ws[ws.max_row]:
            cell.fill = _SUBASSEMBLY_FILL
            cell.font = _SUBASSEMBLY_FONT
        for row in children:
            values = _row_values(item_no, row)
            values[1] = "    " + str(values[1])  # indent under its sub-assembly
            ws.append(values)
            item_no += 1

    _autosize(ws)
    _append_totals(ws, bom_result)
    wb.save(output_path)
    return output_path


def export_bom(bom_result: dict, output_path: str, indented: bool) -> str:
    """Write `bom_result` (a CadAdapter.get_assembly_bom() return value) to
    `output_path` as an .xlsx file, flat or indented per `indented`.
    Returns `output_path` on success.
    """
    if indented:
        return export_bom_indented(bom_result, output_path)
    return export_bom_flat(bom_result, output_path)
